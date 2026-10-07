#!/usr/bin/env python3
"""Build org-chart.html — interactive hierarchy view of the agent org map.

Reads agents_data.py (roster + reporting lines) and the recalculated workbook
for responsibility counts. DERIVATION, not canon.
"""
import json, os, sys
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from agents_data import A, DEPTS, DOTTED, UNITS, AUTONOMY

wb = load_workbook(os.path.join(HERE, "GFunnel-Agent-Org-Map.xlsx"), data_only=True)
counts = {r[0]: r[19] for r in wb["Agent Roster"].iter_rows(min_row=2, values_only=True) if r[0]}
if any(v is None for v in counts.values()):
    sys.exit("Workbook has no cached values — recalculate it first.")
owned = {}
for r in wb["Deliverables"].iter_rows(min_row=5, values_only=True):
    if r[0]:
        owned.setdefault(r[2], []).append(f"{r[1]} — team: {r[5]}")
QUAD = {"STR":"yin","FIN":"yin","HRC":"yin","CON":"yin","MKT":"yang","SAL":"yang","OPS":"yang","TEC":"bridge","CS":"bridge","GOV":"gov","ORC":"orc"}
agents = []
for a in A:
    agents.append(dict(id=a[0], name=a[1], tier=a[2], dept=a[3], unit=UNITS.get(a[0]) or UNITS.get(a[3]) or DEPTS[a[3]][0],
        kingdom=DEPTS[a[3]][1], yy=DEPTS[a[3]][2], reports=a[4], dotted=DOTTED.get(a[0], "—"), mission=a[5], algs=a[6],
        triggers=a[7], outputs=a[8], handoff=a[9], tools=a[10], kpis=a[11], gate=a[12], guard=a[13], auto=a[14],
        basis=a[15], source=a[16], resp=counts[a[0]], quad=QUAD[a[3]], owns=owned.get(a[0], [])))
data = json.dumps({"agents": agents, "autonomy": AUTONOMY, "total_resp": sum(counts.values())}, ensure_ascii=False)
html = open(os.path.join(HERE, "org_chart_template.html"), encoding="utf-8").read().replace("/*__DATA__*/null", data)
open(os.path.join(HERE, "org-chart.html"), "w", encoding="utf-8").write(html)
print("wrote org-chart.html", len(agents), "agents")
