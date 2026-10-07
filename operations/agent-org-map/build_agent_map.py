#!/usr/bin/env python3
"""Build GFunnel-Agent-Org-Map.xlsx — task-specific agents for a business operation.

DERIVATION, not canon. Reads the canonical v5.1 text directly so every algorithm
step becomes an owned responsibility; nothing is retyped by hand except the
operational tables (pipeline, script, I.A.C.E., Immersion, BEAS, Five Modes...).

    python3 operations/agent-org-map/build_agent_map.py
"""
import re, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)
from agents_data import A as AGENTS, DEPTS, AUTONOMY, DOTTED, UNITS, TEAMS

CANON = os.path.join(ROOT, "versions/v5.1/GFunnel-Methodology-v5.1.md")
OUT = os.path.join(HERE, "GFunnel-Agent-Org-Map.xlsx")

# ── 1. Parse every algorithm in v5.1 (Master Meta → Ecological Restoration) ──
def parse_algorithms():
    L = open(CANON, encoding="utf-8").read().split("\n")
    algs, cur = [], None
    for i, line in enumerate(L, 1):
        if i < 2169 or i > 4002:
            continue
        m = re.match(r"^#{2,4} (.+)$", line)
        if m:
            t = m.group(1).strip()
            if cur and cur["steps"]:
                algs.append(cur)
            if t == "The 12 Steps":
                t = "Master Meta-Algorithm"
            cur = {"name": t, "start": i, "steps": []} if ("Algorithm" in t and "Index" not in t and "Per-Layer" not in t) else None
            continue
        if cur:
            s = re.match(r"^(\d+)\.\s+(.*)$", line)
            if s:
                cur["steps"].append((int(s.group(1)), s.group(2).strip(), i))
    if cur and cur["steps"]:
        algs.append(cur)
    return algs

ALGS = parse_algorithms()

# ── 2. Algorithm → owner, category, coverage status, cadence ──
# owner may be a str (whole algorithm) or {step: owner} with "*" default.
OWN = {
 "Master Meta-Algorithm": ("O-02","Meta","Assigned","Per un-routed problem"),
 "Layer 0 —": ("O-01","Per-layer","Assigned","Per input"),
 "Layer I —": ("O-05","Per-layer","Assigned","Per diagnosis"),
 "Layer I.B": ("O-11","Per-layer","Assigned","Before each scale move"),
 "Layer I.C": ("O-12","Per-layer","Assigned","Per unfamiliar process"),
 "Layer I.D": ("O-07","Per-layer","Assigned","Continuous / monthly"),
 "Layer I.E": ("O-08","Per-layer","Assigned","Per stall"),
 "Layer I.F": ("O-09","Per-layer","Assigned","Per asset that must persist"),
 "Layer I.G": ("O-10","Per-layer","Assigned","Per failure"),
 "Layer II —": ("O-06","Per-layer","Assigned","Per failure diagnosis"),
 "Layer III —": ("O-13","Per-layer","Assigned","Per system design review"),
 "Layer III+": ("O-14","Per-layer","Assigned","Per cycle kickoff"),
 "Layer IV": ({"*":"STR-01","1":"G-00","2":"G-00","5":"G-00","7":"CON-00"},"Per-layer","Assigned (founder decides steps 1, 2, 5)","Annual / quarterly"),
 "Layer V —": ({"*":"O-04","6":"O-03"},"Per-layer","Assigned","Monthly"),
 "Layer VI —": ({"*":"O-03","6":"O-15"},"Per-layer","Assigned","Per build / fix / extend"),
 "Layer VII —": ({"1":"MKT-00","2":"SAL-05","3":"SAL-06","4":"CS-09","5":"SAL-03","6":"SAL-12","7":"SAL-00"},"Per-layer","Assigned","Continuous"),
 "Layer ◇": ("G-02","Per-layer","Assigned","Per claim"),
 "Negotiation": ("SAL-10","Business operations","Assigned","Per negotiation"),
 "Software Development": ("TEC-07","Business operations","Assigned","Per software build"),
 "Sales Algorithm": ({"1":"MKT-03","2":"SAL-01","3":"SAL-05","4":"SAL-05","5":"SAL-05","6":"SAL-05","7":"SAL-05","8":"SAL-07","9":"SAL-06","10":"SAL-08","11":"SAL-11","12":"CS-09"},"Business operations","Assigned","Per lead"),
 "Hiring": ({"*":"HRC-01","8":"HRC-02","9":"HRC-03","10":"HRC-04"},"Business operations","Assigned (hire decision human)","Per role"),
 "Client Onboarding": ({"1":"CS-01","2":"CS-02","3":"CS-02","4":"CS-03","5":"CS-03","6":"OPS-05","7":"CS-04","8":"CS-06","9":"CS-10","10":"O-15"},"Business operations","Assigned","Per new client"),
 "Decision-Making": ({"*":"O-16","8":"G-00"},"Business operations","Assigned (choice human)","Per material decision"),
 "Conflict Resolution": ("HRC-08","Business operations","Assigned","Per conflict"),
 "Strategic Planning": ("STR-03","Business operations","Assigned","Quarterly; monthly review"),
 "Personal Development": ("G-07","Life domain → operator","Assigned (operator support, human-led)","Weekly"),
 "Relationships": ("G-00","Life domain","Human-reserved","—"),
 "Health and Body": ("G-00","Life domain","Human-reserved","—"),
 "Financial Decisions": ("FIN-01","Life domain → business capital","Assigned","Per capital decision; monthly review"),
 "Scientific Inquiry": ("O-22","Scientific & cognitive","Assigned","Per Acquire-phase question"),
 "Learning Algorithm": ("HRC-07","Scientific & cognitive","Assigned","Per learning plan"),
 "Problem-Solving": ("O-17","Scientific & cognitive","Assigned","Per obstacle"),
 "Critical Thinking": ("O-18","Scientific & cognitive","Assigned","Per relied-on claim"),
 "Memory and Recall": ("CON-08","Scientific & cognitive","Assigned","Quarterly review"),
 "Creative Work": ("CON-01","Creative & communication","Assigned","Per asset"),
 "Writing Algorithm": ({"1":"CON-T1","2":"CON-T1","3":"CON-09","4":"CON-10","5":"CON-10","6":"CON-10","7":"CON-11","8":"CON-11","9":"CON-T1"},"Creative & communication","Assigned (split across the Editorial & Email team)","Per writing task"),
 "Communication and Speaking": ("MKT-06","Creative & communication","Assigned","Per communication"),
 "Teaching and Mentoring": ("CON-05","Creative & communication","Assigned","Per course / mentoring cycle"),
 "Domain 1 —": ("G-00","Universal process domain","Human-reserved","—"),
 "Domain 2 —": ("G-00","Universal process domain","Human-reserved","—"),
 "Domain 3 —": ("G-00","Universal process domain","Human-reserved (psychological care is not delegated to agents)","—"),
 "Domain 4 —": ("O-01","Universal process domain","Assigned (information overwhelm triage)","Per input surge"),
 "Domain 5 —": ("O-22","Universal process domain","Assigned (formal analysis)","Per analysis"),
 "Domain 6 —": ("CON-06","Universal process domain","Assigned (precise communication)","Per writing task"),
 "Domain 7 —": ("STR-04","Universal process domain","Assigned","Quarterly"),
 "Domain 8 —": ("O-12","Universal process domain","Reference library (Correspondence lookup)","Per mapping"),
 "Domain 9 —": ("O-12","Universal process domain","Reference library (Correspondence lookup)","Per mapping"),
 "Domain 10 —": ("O-12","Universal process domain","Reference library (Correspondence lookup)","Per mapping"),
 "Project Management": ("OPS-03","Extended business","Assigned","Per project; weekly cycles"),
 "Crisis Response": ({"*":"O-19","3":"G-00"},"Extended business","Assigned (command locus human)","Per crisis"),
 "Brand and Identity": ("CON-04","Extended business","Assigned","Per brand review"),
 "Customer Service": ("CS-05","Extended business","Assigned","Per ticket"),
 "Vendor Management": ("OPS-04","Extended business","Assigned","Per vendor; monthly review"),
 "Innovation and R&D": ("STR-06","Extended business","Assigned","Per innovation cycle"),
 "Risk Management": ("STR-05","Extended business","Assigned","Monthly"),
 "Compliance and Legal": ("HRC-10","Extended business","Assigned","Quarterly minimum"),
 "System Architecture": ({"*":"TEC-01","8":"TEC-08"},"Extended business","Assigned","Per system; each Fibonacci step"),
 "Parenting": ("G-00","Extended life","Human-reserved","—"),
 "Time Management": ("G-07","Extended life → operator","Assigned (operator support, human-led)","Weekly"),
 "Energy Management": ("G-07","Extended life → operator","Assigned (operator support, human-led)","Monthly"),
 "Habit Formation": ("G-07","Extended life → operator","Assigned (operator support, human-led)","Daily tracking"),
 "Skill Acquisition": ("HRC-07","Extended life → team","Assigned","Per skill plan"),
 "Grief": ("G-00","Extended life","Human-reserved","—"),
 "Major Life Transition": ("G-00","Extended life","Human-reserved","—"),
 "Pattern Recognition": ("O-20","Extended cognitive","Assigned","Monthly"),
 "Synthesis": ("O-21","Extended cognitive","Assigned","Per synthesis"),
 "Forecasting": ("FIN-03","Extended cognitive","Assigned","Monthly"),
 "Investigation and Research": ("O-22","Extended cognitive","Assigned","Per question"),
 "Public Speaking": ("MKT-06","Extended communication","Assigned (delivery human)","Per talk"),
 "Group Facilitation": ("HRC-09","Extended communication","Assigned","Per meeting"),
 "Influence": ("SAL-10","Extended communication","Assigned","Per influence goal"),
 "Meditation Practice": ("G-00","Extended contemplative","Human-reserved","—"),
 "Ritual Design": ("HRC-05","Extended contemplative → culture","Assigned (team rituals)","Per ritual"),
 "Governance Design": ("G-06","Extended governance","Assigned","Multi-year; on structural change"),
 "AI Alignment": ("G-05","Extended governance","Assigned (applied to this agent fleet)","Every capability scale-up"),
 "Movement Building": ("CS-11","Extended governance","Assigned","Per cycle"),
 "Ecological Restoration": ("G-00","Extended governance","Out of scope for a non-ecological business (human-reserved)","—"),
}
def lookup(name):
    hits = [k for k in OWN if name.startswith(k)] or [k for k in OWN if k in name]
    if len(hits) == 1:
        return OWN[hits[0]]
    raise KeyError(f"{name}: {hits}")

# ── 3. Hand-transcribed operational tables → responsibilities ──
# (element_key, agent, responsibility, source, canon lines, cadence, trigger, output, handoff, human_gate)
T = []
def t(*r): T.append(r)
DEPT_FUNCS = {
 "STR":[("Vision","STR-01"),("Mission","STR-01"),("Values","STR-02"),("Governing principles","STR-02"),("Quarterly direction","STR-03"),("Competitive positioning","STR-04")],
 "MKT":[("Lead generation","MKT-01"),("Traffic","MKT-02"),("Attention","MKT-03"),("Brand awareness","MKT-04"),("Content distribution","MKT-05")],
 "SAL":[("Offer architecture","SAL-07"),("Pipeline","SAL-03"),("Sales scripts","SAL-05"),("ACE System","SAL-00"),("Revenue generation","SAL-00")],
 "OPS":[("Fulfillment","OPS-01"),("Delivery","OPS-01"),("Quality control","OPS-02"),("Project management","OPS-03")],
 "FIN":[("Capital management","FIN-01"),("Budgeting","FIN-02"),("Forecasting","FIN-03"),("Reporting","FIN-04"),("Commission","FIN-05"),("Pricing","FIN-06")],
 "HRC":[("Hiring","HRC-01"),("Onboarding","HRC-03"),("Performance","HRC-04"),("Culture","HRC-05"),("Compensation","HRC-02"),("Org design","HRC-06")],
 "TEC":[("Platform architecture","TEC-01"),("Automation (n8n)","TEC-02"),("CRM (GHL)","TEC-03"),("Integrations","TEC-04"),("AI","TEC-05"),("Security","TEC-06")],
 "CS":[("Onboarding","CS-01"),("Support","CS-05"),("Account management","CS-04"),("NPS","CS-07"),("Churn prevention","CS-08"),("Expansion","CS-09")],
 "CON":[("IP creation","CON-01"),("Media production","CON-02"),("Methodology documentation","CON-03"),("Brand","CON-04"),("GFunnel University","CON-05")],
}
LINE = {"STR":1378,"MKT":1379,"SAL":1380,"OPS":1381,"FIN":1382,"HRC":1383,"TEC":1384,"CS":1385,"CON":1386}
for d, fns in DEPT_FUNCS.items():
    for fn, owner in fns:
        t(f"DEPT:{d}:{fn}", owner, f"Own the '{fn}' function of the {DEPTS[d][0]} ({DEPTS[d][3]}).", "Nine Departments map — primary function", f"v5.1 L{LINE[d]}", "Continuous", "—", "Function performing; evidence for hub BEAS row", d+"-00", "")
    t(f"DEPT:{d}:BEAS-row", d+"-00", f"Self-score the {DEPTS[d][0]} on Correctness, Complexity, Patience, Resilience (0 / 0.5 / 1 each) with evidence, and score coordination with the other eight hubs (0–1).", "BEAS Scoring Instructions", "v5.1 L1396–1404", "Monthly", "BEAS cycle", "Hub row submitted to O-04", "O-04", "Evidence sign-off")

SW = [("01","Direct","Define before you gather — a single clear statement of intent, direction, scope.","O-03","Begin gathering without a clear direction",1488,1497,"Direct sign-off for strategic work"),
      ("02","Guide","Structure the intake — departments in scope, questions, stakeholders, systems to audit — before information arrives.","O-03","Building the container after the information arrives (5× cost)",1500,1509,""),
      ("03","Gather","Blueprint first — collect everything (SOPs, tribal knowledge, tools, pain points, workflows, expectations, revenue flows, contractor relationships); filter nothing.","O-03","Filtering during gathering",1512,1521,""),
      ("04","Organize","Socks before shoes — categorize, find patterns, map dependencies, sequence; produce blueprint, dependency map, prioritized gap list, build sequence.","O-03","Creating before organizing (3–10× rework)",1524,1533,""),
      ("05","Create","Knowledge is free, you pay for creation — build only from the organized blueprint.","O-03","Skipping steps 01–04: the wrong thing perfectly built",1536,1545,""),
      ("06","Database","If it isn't written, it isn't real — SOPs, Komodo recordings, transcripts, same-day client email recap.","O-15","Skipping documentation because the work is done",1548,1557,""),
      ("07","Repetition","Incremental correctness compounds — restart from the higher baseline; New Thought re-enters at Gather.","O-03","Stopping at Create without 06–07",1560,1569,"")]
for n,nm,txt,own,fail,a,b,g in SW:
    t(f"SW:{n}", own, f"Shepherd's Way Step {n} — {nm}: {txt} Guard against: {fail}.", "Shepherd's Way step table", f"v5.1 L{a}–{b}", "Per work item", "Previous step complete", f"Step {n} output", "Next step owner", g)
t("SW:NewThought","O-03","New Thought Principle: route any new insight/input back into Gather (not the beginning); the cycle incorporates without restarting.","New Thought Principle","v5.1 L1468–1471","Per new input","New insight mid-cycle","Item re-entered at Gather","Current cycle owner","")

IM = [("PreImmersion","Pre-Immersion Prep (1–2 days before): review all client materials; identify known knowns, unknowns, gaps; build intake structure → structured framework, variable list, questions.","CS-01",1335),
      ("Calibration","Arrival & Calibration (first 30 min): establish rapport, frame the day as data collection, match the client's frequency → psychological safety.","CS-02",1336),
      ("FullCapture","Full Data Capture (6–12 h): record everything on Komodo — workflows, tribal knowledge, team dynamics, pain points, tools, revenue flows, customer journey, founder vision → raw complete data set, no interpretation.","CS-02",1337),
      ("Organize","Organize Phase (same day, final 2 h): categorize by department, identify patterns, map dependencies, flag contradictions → organized blueprint, prioritized gap list.","CS-03",1338),
      ("Blueprint","Blueprint Delivery (within 48 h): current state map, gaps, Fibonacci build sequence, BEAS baseline, 90-day roadmap → written, client-approved blueprint.","CS-03",1339),
      ("Build","Build Phase (per roadmap): execute against blueprint; every session recorded; every deliverable emailed same day → documented, reproducible system.","OPS-05",1340)]
for k,txt,own,l in IM:
    t(f"IMM:{k}", own, "Immersion Model — "+txt, "Immersion Model table", f"v5.1 L{l}", "Per new client", "—", "Phase output", "Next phase owner", "Client approval" if k=="Blueprint" else ("Session human-led" if k in("Calibration","FullCapture") else ""))

for k,own,txt in [("Centralize","MKT-03","Centralize all lead channels into Lead Connector (GHL)."),
                  ("Tag","MKT-03","Tag leads with source, campaign and hot buttons at capture."),
                  ("InstantReply","SAL-01","Flows AI replies instantly to all inbound — zero lag."),
                  ("Qualify","SAL-02","Qualify quickly through automated short surveys."),
                  ("ChannelROI","MKT-03","Track channel ROI at capture — reallocate to highest-converting."),
                  ("YinYang","MKT-00","Run automated inbound (Yin) and active outreach (Yang) simultaneously; avoid outreach without qualification and inbound with no outreach.")]:
    t(f"ACE:A:{k}", own, "ACE Acquisition — "+txt, "ACE Stage 01 core actions", "v5.1 L1593–1597", "Continuous", "New lead / weekly review", "—", "SAL-03", "Budget reallocation" if k=="ChannelROI" else "")
for st,own in [("New Lead","MKT-03"),("Contacted","SAL-01"),("Qualified","SAL-02"),("Appointment Set","SAL-04"),("Show","SAL-04"),("Proposal","SAL-08"),("Closed Won","SAL-11"),("Closed Lost","O-10")]:
    t(f"PIPE:{st}", own, f"Pipeline stage '{st}': own entry to this stage and its automated follow-up sequence so no lead falls through without a trigger." + (" Closed Lost → integrate the lost path (Layer I.G) and re-nurture." if st=="Closed Lost" else ""), "ACE GHL pipeline", "v5.1 L1596", "Per lead", "Stage entered", "Stage automation fired", "SAL-03", "")
for k,own,txt in [("Script","SAL-05","Follow the 7-step script structure."),("RecordTag","SAL-09","Record and tag all calls in Lead Connector for QA."),("AIContracts","SAL-08","AI populates contracts and proposals from call notes."),("Stack","SAL-07","Stack three offers: Software / Done-For-You / Custom."),("FuturePace","SAL-05","Future pace before every major transition."),("YinYang","SAL-00","Hold each call at the middle: enough discovery (Yin) for the offer to land, enough confidence (Yang) to ask for the decision.")]:
    t(f"ACE:C:{k}", own, "ACE Creation — "+txt, "ACE Stage 02 core actions", "v5.1 L1606–1609", "Per call", "Call booked/held", "—", "SAL-11", "Call is human-led" if k in("Script","FuturePace","YinYang") else "")
for tier,txt in [("1","Software Only: access to the GFunnel platform."),("2","Done-For-You: GFunnel builds and manages."),("3","Custom Enterprise: full white-label or custom build.")]:
    t(f"OFFER:T{tier}","SAL-07",f"Maintain Offer Tier {tier} — {txt} Every prospect has a landing point.","ACE Offer Architecture","v5.1 L1609","Quarterly review","—","Tier definition","FIN-06","Offer changes")
SC=[("01","Permission + Frame","Yin","Establish psychological safety and mutual agenda.","Are you open to me asking some questions before we dive in?"),
    ("02","Discovery","Yin — pure listening","Understand the prospect's actual situation.","Tell me about your current situation. What's working? What's not?"),
    ("03","Pain Excavation","Yin→2nd","Quantify the cost of the current situation.","What does that cost you? In time? In revenue? In stress? What happens if nothing changes?"),
    ("04","Vision Pull","2nd→Yang","Establish the desired state. The gap IS the sale.","If we could solve [specific problem], what does that make possible for you?"),
    ("05","Solution Bridge","Yang — presenting","Connect their pain/vision to the specific solution.","Here's exactly how we solve that: [specific system to specific need]."),
    ("06","Offer + Stack","Yang — the ask","Present all three tiers. Let them self-select.","We have three ways to engage. Let me walk you through each."),
    ("07","Close + Handoff","Yang closing, Yin receiving","Ask for the decision; handle objections via I.A.C.E.; transition to onboarding.","Based on everything you've shared, which option makes the most sense for where you are right now?")]
for n,nm,yy,p,ph in SC:
    t(f"SCRIPT:{n}","SAL-05" if n!="07" else "SAL-11" if False else "SAL-05",f"7-Step Script {n} — {nm} ({yy}): {p} Prompt: \"{ph}\"","7-Step Sales Script","v5.1 L1615–1623","Per call","Previous script step","Prospect moved to next step","SAL-06 (objections) / SAL-11 (handoff)","Spoken by human closer")
for L_,nm,txt in [("I","Isolate","Identify the real objection beneath the surface objection."),("A","Acknowledge","Match frequency; validate the concern (Vibration: match before you shift)."),("C","Clarify","Get precise — what exactly is the variable?"),("E","Eliminate or Escalate","Resolve with specific evidence, or acknowledge a legitimate reason not to proceed (the honest path).")]:
    t(f"IACE:{L_}","SAL-06",f"I.A.C.E. {L_} — {nm}: {txt}","I.A.C.E. Objection Loop","v5.1 L1632–1637","Per objection","Objection raised","Objection resolved or honestly escalated","SAL-05","Escalate decision by human closer")
for k,own,txt in [("Referral","CS-09","Flows AI automates referral sequences after successful deliverables."),("ReportAuto","CS-06","Measure and report results automatically."),("HourBonus","CS-09","Hour bonuses trigger natural upgrade conversations."),("Downgrade","CS-08","Downgrade path to $297 prevents full cancellation."),("Community","CS-11","Community + social network creates stickiness beyond product."),
                  ("MonthlyBEAS","CS-06","Monthly BEAS review delivered to client."),("QuarterlyImm","CS-10","Quarterly Immersion review session."),("ProgressReport","CS-06","Progress report showing First Kingdom results from Second Kingdom systems built — the report IS the upsell trigger."),("YinYang","CS-00","Hold expansion at the middle: proactive result reporting creates natural expansion; neither constant upsell nor serve-and-wait.")]:
    t(f"ACE:E:{k}", own, "ACE Expansion — "+txt, "ACE Stage 03 core actions / retention architecture", "v5.1 L1646–1650", "Per client cycle", "—", "—", "CS-00", "")
PILL=[("Correctness","Every move objectively sound; no shortcuts (97% accuracy → 16% perfection over 60 decisions). In practice: documented SOPs, tested automations before launch, blueprint review before any build, every Komodo recording.","OPS-02",1364),
      ("Complexity","Embrace full scope; never oversimplify. In practice: nine integrated departments, no isolated tools, Gather step before any build.","TEC-04",1365),
      ("Patience","Growth through position, not desperation; blueprint first. In practice: 5-phase implementation, Direct → Guide → Gather before any system is built, quarterly BEAS assessment.","O-03",1366),
      ("Resilience","Survive pressure without collapse; productive entropy. In practice: Komodo-recorded processes, full documentation, never fully dependent on any single third party, downgrade path to $297.","STR-05",1367)]
for p,txt,own,l in PILL:
    t(f"PILLAR:{p}",own,f"Pillar steward — {p}: {txt}","Four Pillars table",f"v5.1 L{l}","Continuous; scored monthly","—","Pillar evidence across all hubs","O-04","")
for k,own,txt in [("Score","O-04","Score each of 36 cells 0 / 0.5 / 1 (absent / partial / fully implemented)."),("DeptTotal","O-04","Department Total = sum of four pillar scores (max 4)."),("Coord","O-04","Coordination Score 0–1 per department (max 9)."),("Total","O-04","BEAS Total = 36 pillar scores + coordination (max 45)."),
                  ("Band40","O-04","40–45 Dynamic Middle—Sustained → scale intentionally; begin next Fibonacci step."),("Band32","O-04","32–39 Equilibrium—Functional → identify lowest department; apply Shepherd's Way."),("Band24","O-04","24–31 Imbalance—Managed → gap analysis; prioritize by dependency order."),("Band16","O-04","16–23 Imbalance—Exposed → return to Gather; full Immersion Model recommended."),("Band0","O-19","0–15 Collapse Risk → emergency blueprint session; minimum viable docs in all 9 departments within 30 days."),
                  ("QYang","O-07","Quadrant check — High Yang (Marketing, Sales, Operations): overdeveloped = revenue without infrastructure; underdeveloped = no pipeline/clients/capacity."),("QYin","O-07","Quadrant check — High Yin (Strategy, Finance, HR & Culture, Content): overdeveloped = organized business that doesn't grow; underdeveloped = no direction/capital control/team/IP."),("QBridge","O-07","Quadrant check — Bridging (Technology, Client Success): overdeveloped = over-engineered systems nobody uses; underdeveloped = manual processes at scale, churn.")]:
    t(f"BEAS:{k}",own,"BEAS — "+txt,"BEAS scoring / interpretation / quadrant tables","v5.1 L1389–1446","Monthly","BEAS cycle","—","O-03","Stage classification" if k.startswith("Q") else "")
for m,nm,own,txt in [("1","Direct Teaching","CON-05","Person-to-person transfer — sales calls, onboarding, coaching."),("2","Demonstration","CON-02","Show the process in action — Komodo recordings, case studies, testimonials."),("3","Documentation","CON-03","Encode into text/video/structured media — SOPs, methodology, GFunnel University."),("4","System-Building","TEC-02","Systems that execute the process automatically — n8n automations, GHL pipelines, AI workflows."),("5","Cultural Embedding","HRC-05","The process becomes a community assumption — community norms, brand language, methodology.")]:
    t(f"MODE:{m}",own,f"Propagation Mode {m} — {nm}: {txt}","Five Modes of Propagation","v5.1 L1303–1309","Continuous","—","Process independent of its originator","CON-00","")
for k,own,txt in [("Docs","O-15","Sustain: documentation at every step / Kill: tribal knowledge."),("Results","CON-07","Sustain: results that speak / Kill: promises without proof."),("Community","CS-11","Sustain: community formation (Mode 5) / Kill: transactional relationships."),("SysIndep","TEC-02","Sustain: system independence via Mode 4 / Kill: key-person dependencies."),("Fibonacci","O-11","Sustain: Fibonacci growth, each cycle informed by the prior two / Kill: exponential growth without Yin infrastructure (Texas Grid).")]:
    t(f"PROP:{k}",own,"Propagation guard — "+txt,"What Sustains vs. What Kills Propagation","v5.1 L1318–1324","Monthly audit","—","—","CON-00","")
for n,txt in enumerate(["Identified which algorithm applies (else Master Meta-Algorithm / Algorithm Index).","Ran the algorithm in order, without skipping steps.","Held open the variables named as held open.","Documented the run (Layer VI Step 6).","Applied Correspondence at adjacent scales.","Distinguished measured from assumed (Variable Principle).","Integrated failed paths from prior runs (Layer I.G)."],1):
    t(f"AUDIT:{n}","G-04",f"Self-audit check {n}: {txt}","Self-Audit Checklist","v5.1 L182–203","Per agent run","Run completed","Pass/fail","Failing agent","")
for n,txt in enumerate(["The choice of action at any decision point.","The naming of outcome after the action.","Empirical measurements specific to the situation.","Stance on bounded variables the framework holds open."],1):
    t(f"CANNOT:{n}","G-00",f"Supply what the framework cannot: {txt}","What The Framework Cannot Supply","v5.1 L168–181","Per gate","Human Gate raised","Decision / name / measurement","Requesting agent","Human-only")
for n,txt in enumerate(["Step 1 skipped (no domain) → noise.","Step 4 skipped (no desired state) → no target.","Step 5 skipped (no dynamic middle) → burnout or paralysis.","Step 9 sequencing skipped (Third → First without Second) → wrong thing perfectly built.","Step 11 skipped (no documentation) → dependency, dies with operator.","Step 12 skipped (no higher baseline) → flat repetition."],1):
    t(f"MMFAIL:{n}","G-04",f"Detect Master Meta-Algorithm failure mode {n}: {txt}","Master Meta-Algorithm Failure Modes","v5.1 L2206–2225","Per derivation","Derived run submitted","Failure flagged","O-02","")
for L_,txt in [("A","Name & Reframe"),("B","Structure (DETECT / PROCESS / RESPOND)"),("C","Placement (Kingdom, domain, principle)"),("D","Gradient & Dynamic Middle"),("E","Forced vs Guided"),("F","Keys & Doors — Decode / Drown / Numb"),("G","Boundary Class — Scaffolding vs Constitutive"),("H","Failure Mode & Rebuild"),("I","Falsifiability & Variable Ledger"),("J","Leverage — Knowledge → Action → Discipline"),("K","Correspondence Check"),("L","Document & Restart")]:
    t(f"DEEPLENS:{L_}","O-02",f"Deep-Lens {L_} — {txt}: run on any business mechanism needing a high-resolution diagnosis.","v5.2 Deep-Lens Protocol §1","v5.2 L29–74","Per deep diagnosis","Deep diagnosis requested","Deep-Lens record","G-01","")
for k,txt in [("A","Test A · Ground — is the gap within-system or a source/ground question?"),("B","Test B · Uniqueness — does exactly one candidate survive?"),("C","Test C · Direction — what does the structure operate on?"),("D","Test D · Falsifiability — state and execute the falsifier."),("OUT","Output the class: forced-fill · live hypothesis · stipulation · proven-open."),("ANTI","Anti-Operation — name the gap each fill opens (its conjugate); log it."),("ZERO","Perfection Criterion — zero hidden gaps; open centre stays proven-open, not sealed.")]:
    t(f"FT:{k}","G-01","Forcing Test / Container — "+txt,"v5.3 Meta-Tier §1–§3","v5.3 §1–§3","Per open question","Variable proposed for resolution","Classified gap + next gap","G-03","Stipulations → G-00")
for k,txt in [("Measured","Label measured / empirically settled values."),("Derived","Label structurally derived conclusions."),("Open","Label held-open variables; never fill with assumption."),("Narrow","Log each narrowing: what narrowed, what stayed open, what newly opened.")]:
    t(f"VAR:{k}","G-03","Variable Register — "+txt,"Variable Principle / Registers","v5.1 L223–263","Continuous","New value or assumption","Register row","G-01","Measurements → G-00")
t("UPDATE:Protocol","G-01","Document Update Protocol: when evidence narrows a business variable, record the narrowing, the newly opened variables, and load it as substrate for the next cycle (never silently overwrite).","Document Update Protocol","v5.1 L204–222","Per narrowing","Evidence arrives","Iteration entry","G-03","")

# ── 3b. Team-lead responsibilities (DERIVED: Shepherd's Way at team scale) ──
TEAM_DUTIES = [
 ("Direct", "Accept each deliverable request and state its outcome, owner and done-criterion in one sentence before any work starts.", "v5.1 L1488–1497"),
 ("Guide", "Brief and assign: split the deliverable across team members by role and set the intake (inputs, sources, deadlines).", "v5.1 L1500–1509"),
 ("Organize", "Sequence the team's work by dependency (socks before shoes) and keep the team queue visible.", "v5.1 L1524–1533"),
 ("Approve", "Check every deliverable against its done-criterion before it leaves the team (Correctness pillar).", "v5.1 L1364"),
 ("Database", "Confirm each finished item is documented (SOP, recording, recap) before it is closed.", "v5.1 L1548–1557"),
 ("Repetition", "Run a short retro per cycle: integrate failures (Layer I.G) and raise the team's baseline.", "v5.1 L1560–1569; L2374–2391"),
 ("Escalate", "Escalate anything above the team's autonomy, and every Human Gate, to the hub lead or G-00.", "Human Gates tab"),
 ("BEAS evidence", "Supply the team's evidence for its hub's four-pillar BEAS row each month.", "v5.1 L1396–1404"),
]
TEAM_R = []
for tid, (tname, code, boss, mission, members, algs) in TEAMS.items():
    for k, txt, ln in TEAM_DUTIES:
        TEAM_R.append((tid, f"TEAM:{tid}:{k}", f"Team lead — {k}: {txt} Team: {', '.join(members)}.", "Shepherd's Way at team scale (derived)", ln, "Per deliverable / weekly", "Deliverable request" if k in ("Direct","Guide") else "Team cycle", f"{k} complete", boss if k in ("Escalate","BEAS evidence") else ", ".join(members), "Above-autonomy items" if k == "Escalate" else ""))
EDITORIAL = [
 ("CON-09","Plan each issue: one reader outcome, sections, sources (proof library, University, research), length.","Writing step 1–3"),
 ("CON-09","Assemble the issue from approved pieces and hand the complete draft to the editor.","Writing step 3"),
 ("CON-10","Maintain the email template set other agents send (instant reply, follow-ups, referral ask, recap) in the brand voice.","Writing steps 4–6"),
 ("CON-10","Write each email around one specific reader action (Influence: be specific about the shift).","Influence steps 1–2"),
 ("CON-06","Write long-form pieces (articles, guides, issue features) from the issue brief.","Writing step 6; Domain 6"),
 ("CON-11","Fact-check names, numbers, links and dates; nothing goes out unverified.","Variable Principle"),
 ("CON-12","Check that every quoted result is measured and every promise is falsifiable; reject unsupported claims.","Layer ◇; Critical Thinking"),
 ("CON-12","Check brand voice and visual conformance against the brand standards (Correspondence across touchpoints).","Brand step 6"),
 ("CON-12","Carry CC BY 4.0 attribution whenever the methodology is quoted or closely paraphrased.","LICENSE; ATTRIBUTION.md"),
 ("CON-13","Keep the editorial calendar on a steady cadence (Rhythm); avoid clashing sends across hubs.","Communication step 7"),
 ("CON-13","Build, segment and schedule sends in Lead Connector; run a test send before every live send.","Pillar I Correctness"),
 ("CON-14","Report each send against its intended action at 48 h and 7 days; name the weakest step.","Creative Work step 9; Layer VII step 6"),
 ("CON-14","Feed recurring patterns to the Pattern Recognition Agent and the next issue brief.","Pattern Recognition"),
]
for ag_, txt, src in EDITORIAL:
    TEAM_R.append((ag_, f"EDITORIAL:{ag_}:{src}", txt, f"Editorial & Email team (derived) — {src}", "v5.1 L3004–3023", "Per issue / send", "Brief or draft received", "Step output", "CON-T1", ""))

OPLAYER = [
 ("O-23","Send the Owner Daily Brief: gates waiting, blocked items, today's priorities, new proposals, crisis flags — one screen.","Communication step 5 (density the receiver can hold)","Daily"),
 ("O-23","Send the Owner Weekly Report: per-hub status, critical-path projects, funnel and channel numbers, decisions made, proposals to approve.","Communication; Layer V","Weekly"),
 ("O-23","Keep the approvals inbox: every Human Gate arrives with its decision brief; present options, the dynamic middle and the falsifier, never a pre-made choice.","Decision-Making steps 1–7; Cannot-Supply item 1","Continuous"),
 ("O-23","Record every Owner decision with its expected outcome in the decision log and relay it to the requesting agent.","Decision-Making step 9","Per decision"),
 ("O-23","Turn each Owner request into a work item via O-01 and confirm receipt the same day.","Layer 0 step 1","Per request"),
 ("O-23","Escalate a confirmed crisis to the Owner immediately, outside the brief cycle.","Crisis Response step 3","Per crisis"),
 ("O-24","Scan signals weekly: BEAS lowest cell, stalled projects (O-08), lowest-conversion funnel step (SAL-12), clients' next Fibonacci step (CS-10), risks (STR-05), patterns (O-20), open variables (G-03).","Layer V step 5; Layer I.E step 8; Layer VII step 6","Weekly"),
 ("O-24","Write each proposal as a card: project, problem, proposed work, owning team lead, expected outcome, falsifier, effort, dependency order, open variables.","Decision-Making steps 7 and 9; Layer ◇","Per proposal"),
 ("O-24","Rank proposals by dependency order (socks before shoes), not by excitement.","Shepherd's Way Step 04; Project Management step 4","Weekly"),
 ("O-24","Never propose growth that outruns infrastructure: check the Four Pillars before any scale proposal.","Layer I.B","Per proposal"),
 ("O-24","After delivery, compare actual vs expected outcome for each approved proposal and send the gap to O-10.","Decision-Making step 10; Layer I.G","Per completed proposal"),
 ("O-25","Keep one board for all work: every item has an owner team lead, a Shepherd's Way step, a due date and blockers.","Layer VI; Project Management step 2","Continuous"),
 ("O-25","Dispatch approved items to the owning team lead; nothing is dispatched straight to a task agent.","Derived (chain of command)","Per item"),
 ("O-25","Collect daily status from team leads and publish the blocked-items list; a blocked item is the location of the work.","How To Run An Algorithm","Daily"),
 ("O-25","Flag items that skipped Shepherd's Way steps (e.g. Create before Organize) to O-03.","Layer VI failure modes","Daily"),
 ("O-25","Feed status and completion data to O-23 for the Owner briefs.","Database step","Daily"),
]
for ag_, txt, src, cad in OPLAYER:
    TEAM_R.append((ag_, f"OPS-LAYER:{ag_}:{src}", txt, f"Operating layer (derived) — {src}", "Derived", cad, "Cadence or event", "Report / item / proposal", "O-23" if ag_ != "O-23" else "G-00", "Owner decides" if ag_ in ("O-23","O-24") else ""))
DELIV = [
 # (deliverable, accountable, doers, algorithm/process, cadence, human gate)
 ("Newsletter issue","CON-T1","CON-09, CON-06, CON-10, CON-11, CON-12, CON-13, CON-14","Writing Algorithm (split across team)","Per editorial calendar","Claims about results → G-00"),
 ("Marketing / nurture email sequence","CON-T1","CON-10, CON-11, CON-12, CON-13, CON-14 (with MKT-T1 for targeting)","Writing; Influence","Per campaign","Full-list sends"),
 ("Automated email templates (instant reply, follow-up, referral ask, recap)","CON-T1","CON-10, CON-11, CON-12 → used by SAL-01, CS-09, O-15","Writing","Quarterly refresh","—"),
 ("Articles, guides and long-form copy","CON-T1","CON-06, CON-11, CON-12","Writing","Per calendar","Publication of claims"),
 ("Same-day client recap email","O-03","O-15 (templates from CON-10)","Shepherd's Way Step 06","Every client session","—"),
 ("Social posts and channel distribution","MKT-T2","MKT-05, MKT-04","Five Modes (Mode 2); Communication","Per schedule","—"),
 ("Webinar or talk","MKT-T2","MKT-06 (script copy from CON-T1)","Public Speaking","Per event","Delivery is human"),
 ("Outbound campaign","MKT-T1","MKT-01, MKT-03","ACE Acquisition (Yang)","Weekly","Targeting / compliance"),
 ("Inbound funnel or landing page","MKT-T1","MKT-02, TEC-T2","ACE Acquisition (Yin); Shepherd's Way","Per campaign","Spend"),
 ("Channel ROI report","MKT-T1","MKT-03","Layer VII step 1","Weekly","Budget reallocation"),
 ("Lead response and qualification","SAL-T1","SAL-01, SAL-02, SAL-04","Sales Algorithm steps 1–2","Per lead","—"),
 ("Clean pipeline (no lead without a trigger)","SAL-T1","SAL-03","ACE GHL pipeline","Daily","—"),
 ("Sales call brief and QA","SAL-T2","SAL-05, SAL-09","7-Step Script","Per call","Call is human-led"),
 ("Objection library","SAL-T2","SAL-06","I.A.C.E.","Continuous","—"),
 ("Proposal and contract","SAL-T3","SAL-08 (legal check HRC-10)","Sales step 10","Per deal","Signature"),
 ("Three-tier offer stack","SAL-T3","SAL-07, FIN-06","ACE Offer Architecture","Quarterly","Offer changes"),
 ("Sales-to-onboarding handoff package","SAL-T3","SAL-11","Sales step 11","Per Closed Won","—"),
 ("Funnel conversion report","SAL-00","SAL-12","Layer VII step 6","Weekly","—"),
 ("Immersion day and 48-hour blueprint","CS-T1","CS-01, CS-02, CS-03","Immersion Model","Per new client","Blueprint approval"),
 ("Quarterly Immersion review","CS-T1","CS-10","Client Onboarding step 9","Quarterly","—"),
 ("Client build per roadmap","OPS-T1","OPS-05, OPS-01, OPS-02","Immersion Build Phase","Per roadmap","—"),
 ("Project plan and weekly cycle","OPS-T2","OPS-03","Project Management","Per project","Scope changes"),
 ("Vendor scorecard and downgrade path","OPS-T2","OPS-04","Vendor Management","Monthly","Contracting"),
 ("Monthly client BEAS and results report","CS-T3","CS-06","ACE Retention Architecture","Monthly","—"),
 ("Referral and upgrade campaign","CS-T3","CS-09 (copy from CON-10)","ACE Expansion","After first result","Upsells above threshold"),
 ("Churn save ($297 downgrade path)","CS-T3","CS-08","ACE Expansion","Per risk signal","Concessions"),
 ("Community programme","CS-T3","CS-11","Movement Building","Monthly","—"),
 ("Support ticket resolution","CS-T2","CS-05","Customer Service","Per ticket","Refunds/credits"),
 ("NPS survey and feedback themes","CS-T2","CS-07","Layer I.F","Quarterly / milestones","—"),
 ("Budget","FIN-T1","FIN-02","Financial Decisions","Quarterly","Approval"),
 ("Forecast","FIN-T1","FIN-03","Forecasting","Monthly","—"),
 ("Monthly financial report","FIN-T1","FIN-04","Database / Layer I.F","Monthly","—"),
 ("Invoices and collections","FIN-T2","FIN-07","Treasury (civil map)","Per invoice","Write-offs"),
 ("Commission statements","FIN-T2","FIN-05","Database","Per pay period","Payouts"),
 ("Price book","FIN-T2","FIN-06","Domain 7","Quarterly","Price changes"),
 ("Capital decision brief","FIN-00","FIN-01","Financial Decisions","Per decision","Capital deployment"),
 ("New hire","HRC-T1","HRC-01, HRC-02, HRC-03","Hiring","Per role","Hire and compensation"),
 ("30/60/90 review","HRC-T2","HRC-04","Hiring step 10","Per hire","Employment decisions"),
 ("Learning plan","HRC-T2","HRC-07","Learning; Skill Acquisition","Per gap","—"),
 ("Meeting agenda and minutes","HRC-T2","HRC-09","Group Facilitation","Per meeting","—"),
 ("Team rituals and culture norms","HRC-T3","HRC-05","Ritual Design","Per ritual","—"),
 ("Org chart and role charters (this map)","HRC-T3","HRC-06","Governance Design","Per growth step","Structure changes"),
 ("Compliance register","HRC-T3","HRC-10","Compliance and Legal","Quarterly","Legal positions"),
 ("Conflict resolution record","HRC-T3","HRC-08","Conflict Resolution","Per conflict","Acceptance by parties"),
 ("Architecture document","TEC-T1","TEC-01","System Architecture","Per system","Major choices"),
 ("Custom software release","TEC-T1","TEC-07, OPS-02","Software Development","Per release","Production deploy"),
 ("n8n automation","TEC-T2","TEC-02","Propagation Mode 4","Per SOP","—"),
 ("CRM configuration change","TEC-T2","TEC-03","ACE GHL architecture","Per request","—"),
 ("New or changed AI agent","TEC-T2","TEC-05 (review by G-05)","AI Alignment","Per change","Deployment"),
 ("Security review","TEC-T3","TEC-06","Risk Management","Quarterly","Sensitive access"),
 ("Monitoring and alerts","TEC-T3","TEC-08","System Architecture step 8","Continuous","—"),
 ("SOP","CON-T3","CON-03 (capture by O-15)","Propagation Mode 3","Per process","—"),
 ("GFunnel University course","CON-T3","CON-05","Teaching and Mentoring","Per course","—"),
 ("Case study or testimonial","CON-T2","CON-02, CON-07","Propagation Mode 2","Per client result","Client consent"),
 ("Brand standards","CON-T4","CON-04","Brand and Identity","Per review","Brand changes"),
 ("New IP (framework, product, course concept)","CON-T4","CON-01","Creative Work","Per plan","Publication"),
 ("Quarterly plan","STR-T1","STR-03, STR-01","Strategic Planning","Quarterly","Approval"),
 ("Values and governing principles","STR-T1","STR-02","Governance Design","Annual","Changes"),
 ("Positioning brief","STR-T2","STR-04","Domain 7","Quarterly","Positioning choice"),
 ("Risk register","STR-T2","STR-05","Risk Management","Monthly","Risk appetite"),
 ("MVP experiment","STR-T2","STR-06","Innovation and R&D","Per cycle","Investment"),
 ("Monthly company BEAS scorecard","O-04","O-07, O-11, O-13 + all hub leads","Layer V","Monthly","Stage sign-off"),
 ("Diagnosis of a stuck process","O-T1","O-05, O-06, O-08, O-09, O-12","Layers I, II, I.E, I.F, I.C","Per request","—"),
 ("Post-mortem / integration record","O-T2","O-10, O-14","Layer I.G","Per failure","—"),
 ("Decision brief","O-T3","O-16, O-17, G-02","Decision-Making","Per decision","The choice"),
 ("Research memo","O-T3","O-22, O-21, O-20","Investigation and Research","Per question","—"),
 ("Crisis log and response","O-19","G-00 (command), MKT-06, O-15","Crisis Response","Per crisis","Command locus"),
 ("Derived algorithm / new SOP for an unnamed problem","O-02","G-01, G-04, CON-03","Master Meta-Algorithm","Per problem","Adoption as SOP"),
 ("Variable register","G-01","G-03","Variable Principle","Continuous","Measurements"),
 ("Agent alignment review","G-01","G-05, G-06","AI Alignment","Per capability change","Deployment"),
 ("Founder weekly time and energy review","G-07","—","Time / Energy Management","Weekly","All personal choices"),
]

DELIV += [
 ("Owner Daily Brief","O-23","O-25 (status), O-24 (proposals), all team leads","Communication; Decision-Making","Daily","—"),
 ("Owner Weekly Report","O-23","All hub leads, SAL-12, MKT-03, OPS-03, O-25","Communication; Layer V","Weekly","—"),
 ("Approvals inbox and decision log","O-23","Every agent raising a Human Gate","Decision-Making steps 8–9","Continuous","Every decision"),
 ("Work proposals (suggested next work)","O-T3","O-24 (signals from O-04, O-08, SAL-12, CS-10, STR-05, O-20)","Layer V step 5; Layer I.E step 8","Weekly","Approval"),
 ("Company work board","O-25","All team leads","Shepherd's Way; Project Management","Continuous","Re-prioritisation"),
]
RHYTHM = [
 # (cadence, report or channel, from, to, contents, anchor)
 ("Continuous","Human Gate request","Any agent","O-23 → Owner","Decision brief: options, variables measured vs open, dynamic middle, falsifier, recommended option","v5.1 L168–181; L2698–2719"),
 ("Continuous","Owner request","Owner","O-23 → O-01 → O-02 → O-25","Request turned into a work item with owner team lead, outcome and due date","Layer 0; Shepherd's Way Step 01"),
 ("Continuous","Crisis alert","O-19","Owner (immediately, via O-23)","Confirmed crisis, impact, triage status, decision needed","v5.1 L3383–3402"),
 ("Per deliverable","Same-day recap","O-15","Client and work board","What was done, value delivered, time spent","v5.1 L1548–1557"),
 ("Daily","Stand-up status","Task agents","Their team lead (on the board)","Done / doing / blocked, current Shepherd's Way step","Derived; v5.1 L96–111"),
 ("Daily","Board update + blocked list","Team leads","O-25","Item states, blockers, items needing a gate","Derived"),
 ("Daily","Owner Daily Brief","O-23","Owner","Gates waiting, blockers, today's priorities, new proposals, crisis flags (one screen)","Derived"),
 ("Weekly","Team report","Team leads","Hub lead","Deliverables shipped, cycle notes, retro and integrated failures","v5.1 L1560–1569"),
 ("Weekly","Hub report","Hub leads","O-23","Hub status, risks, asks of the Owner","Derived"),
 ("Weekly","Funnel and channel numbers","SAL-12, MKT-03","O-23","Conversion per stage, lowest step, channel ROI","v5.1 L2523–2535"),
 ("Weekly","Work proposals digest","O-24 → O-T3","O-23 → Owner","Proposed work with outcome, falsifier, effort; Owner approves / defers / declines","Derived; Layer V step 5"),
 ("Weekly","Owner Weekly Report","O-23","Owner","Per-hub status, projects on critical path, numbers, decisions made, proposals","Derived"),
 ("Weekly","Owner time and energy review","G-07","Owner","Intended vs actual time by Kingdom","v5.1 L3579–3614"),
 ("Monthly","Company BEAS scorecard","O-04","Owner (via O-23)","45-point score, band, quadrant balance, lowest cell and the dispatched fix","v5.1 L2478–2495"),
 ("Monthly","Financial report + forecast","FIN-04, FIN-03","Owner (via O-23)","P&L, cash, forecast vs actual","v5.1 L3749–3768"),
 ("Monthly","Client BEAS + results reports","CS-06","Each client","Results from systems built (the upsell trigger)","v5.1 L1649"),
 ("Monthly","Risk and vendor review","STR-05, OPS-04","Owner (via O-23)","Top risks, downgrade paths, vendor performance","v5.1 L3489–3506; L3445–3462"),
 ("Quarterly","Strategic plan","STR-03","Owner","Plan with assumptions, held-open variables, falsifiers","v5.1 L2744–2765"),
 ("Quarterly","Immersion reviews","CS-10","Each client","Updated blueprint, next Fibonacci step","v5.1 L2693"),
 ("Quarterly","Compliance review","HRC-10","Owner","Compliance register status","v5.1 L3509–3526"),
 ("Per change","Agent alignment review","G-05","Owner","New or expanded agents: purpose, observation, downgrade path","v5.1 L3933–3954"),
]

# ── 4. Workflows ──
WF = [
 ("W01","Lead-to-Cash (ACE Acquisition → Creation)",[("O-01","Input arrives (ad, form, referral)"),("MKT-03","Capture in Lead Connector; tag source/campaign/hot buttons"),("SAL-01","Instant automated first contact (Flows AI) → Contacted"),("SAL-02","Short-survey qualification → Qualified"),("SAL-04","Book and confirm call → Appointment Set / Show"),("SAL-05","Pre-call brief; human closer runs 7-step script"),("SAL-06","I.A.C.E. on objections"),("SAL-07","Present three-tier stack; prospect self-selects"),("SAL-08","Draft proposal/contract from notes → Proposal"),("G-00","Human approves and signs"),("SAL-11","Closed Won → handoff package with full context"),("FIN-07","Invoice and record payment"),("FIN-05","Calculate commission"),("SAL-12","Update step conversion; lowest step → O-03")]),
 ("W02","Client Immersion → Blueprint → Build",[("CS-01","Pre-Immersion prep (1–2 days before)"),("CS-02","Calibration + 6–12 h Komodo capture (human-led)"),("CS-03","Organize same day; blueprint within 48 h"),("G-00","Internal sign-off; client approves blueprint"),("OPS-03","Project plan from 90-day roadmap"),("OPS-05","Build per roadmap; every session recorded"),("OPS-02","QC before handoff"),("O-15","Same-day email recap of every deliverable"),("CS-04","First measurable result within 30 days"),("CS-06","Monthly client BEAS + results report"),("CS-10","Quarterly Immersion review → next Fibonacci step")]),
 ("W03","Expansion & Retention (ACE Expansion)",[("CS-06","Results report delivered (the upsell trigger)"),("CS-09","Referral sequence + hour-bonus upgrade conversation"),("CS-07","NPS / feedback capture"),("CS-08","Churn risk → $297 downgrade path instead of cancellation"),("CS-11","Community engagement for stickiness"),("SAL-05","Upgrade call (human-led) if accepted"),("O-10","Integrate every churn/downgrade as information")]),
 ("W04","Shepherd's Way build cycle (any hub)",[("O-03","01 Direct — intent, scope (strategic: G-00 signs)"),("O-03","02 Guide — intake structure"),("O-03","03 Gather — collect everything, filter nothing"),("O-03","04 Organize — blueprint, dependency map, build sequence"),("Hub agent","05 Create — build from blueprint"),("O-15","06 Database — SOPs, recordings, recap"),("O-03","07 Repetition — restart at higher baseline"),("G-04","Self-audit of the run")]),
 ("W05","Monthly BEAS cycle",[("Hub leads","Submit four-pillar self-scores + coordination with evidence"),("O-04","Total /45, band, quadrant analysis vs growth stage"),("G-00","Confirm growth stage and evidence"),("O-07","Yin/Yang quadrant correction"),("O-04","Dispatch lowest department-pillar cell"),("O-03","Shepherd's Way on that cell"),("O-04","Re-score next month")]),
 ("W06","Hire → Onboard → 30/60/90",[("HRC-01","Role purpose; four-pillar requirements; source; score; Kingdom + frequency check; references"),("G-00","Hire decision"),("HRC-02","Offer at the compensation middle"),("HRC-03","Immersion onboarding; role blueprint; playbook"),("HRC-07","Learning plan"),("HRC-04","30/60/90 reviews with role-fit BEAS"),("HRC-06","Update org design and this agent map")]),
 ("W07","Crisis response",[("O-19","Confirm crisis is real"),("O-19","Stop First-Kingdom bleeding"),("G-00","Take command locus"),("O-19","Triage → stabilization → root cause → prevention"),("MKT-06","Stakeholder communication matched to frequency"),("O-15","Real-time crisis log"),("O-10","Post-crisis integration"),("STR-05","Update risk register and downgrade paths")]),
 ("W08","New Thought re-entry",[("O-01","New insight/input detected mid-cycle"),("O-03","Re-enter at Gather (not the beginning)"),("G-03","Register any new variable it opens"),("O-03","Continue Organize → Create with the new signal")]),
 ("W09","Unknown problem routing",[("O-01","Name input; locate Kingdom; move up"),("O-02","Match in Algorithm Index"),("O-02","No match → run 12-step Master Meta-Algorithm; label DERIVATION"),("G-01","Classify the derived algorithm (forced / live hypothesis / stipulation)"),("G-04","Check failure modes"),("G-00","Adopt as SOP? (human)"),("CON-03","Document derived SOP, labelled as derivation")]),
 ("W10","Failure integration",[("Any agent","Report failure / dead end"),("O-10","Extract information; locate what it enables"),("CON-08","Load into knowledge base (never delete the record)"),("O-14","Load into next cycle's starting substrate"),("O-20","Check for cross-case pattern")]),
 ("W11","Material decision",[("O-16","Decision brief: statement, Kingdom, variables, gradient, Hermetic audit, middle"),("G-02","Attach falsifier"),("G-00","Decide (the choice is human)"),("O-15","Document expected outcome"),("O-16","Compare actual vs predicted"),("O-10","Integrate the gap")]),
 ("W12","Vendor onboarding & review",[("OPS-04","Gradient; Correctness / Complexity / Resilience checks"),("STR-05","Confirm downgrade path"),("HRC-10","Contract review; ambiguous terms held open"),("G-00","Contract approval"),("OPS-04","Instrument performance; monthly review")]),
 ("W13","Content propagation (Five Modes)",[("CON-02","Mode 2 — capture demonstration/case study"),("CON-03","Mode 3 — document as SOP/course material"),("TEC-02","Mode 4 — automate"),("CON-05","Mode 1 — teach via University"),("HRC-05","Mode 5 — embed in culture/community"),("MKT-05","Distribute")]),
 ("W14","Quarterly strategy & compliance review",[("STR-01","Re-filter initiatives against the convergence point"),("STR-04","Market cycle phase and positioning"),("STR-03","Quarterly plan with falsifiers"),("FIN-02","Budget"),("HRC-10","Compliance review (quarterly minimum)"),("G-05","Agent alignment review at any capability change"),("G-00","Approve plan")]),
 ("W16","Newsletter issue (Editorial & Email team)",[("CON-T1","Direct: set the issue's reader outcome and done-criterion"),("CON-09","Plan the issue; pull proof, University and research items; brief writers"),("CON-06","Write feature / long-form sections"),("CON-10","Write emails, subject lines and calls to action"),("CON-11","Subtractive edit; fact-check; read-aloud test"),("CON-12","Brand voice + claims check + attribution"),("CON-T1","Approve"),("CON-13","Build segments, test send, schedule in Lead Connector"),("CON-14","Report at 48 h and 7 days; name the weakest step"),("CON-T1","Document what worked; retro; next issue starts higher")]),
 ("W17","Owner request → done",[("G-00","Owner sends a request (message, email or meeting note)"),("O-23","Acknowledge; clarify the outcome if unclear"),("O-01","Name the input; locate its Kingdom"),("O-02","Route to an algorithm and owning team"),("O-25","Create the board item; dispatch to the team lead"),("Team lead","Direct and Guide; assign the team"),("Task agents","Gather → Organize → Create; daily status on the board"),("Team lead","Approve against the done-criterion"),("O-15","Document; same-day recap"),("O-23","Report completion in the next brief")]),
 ("W18","Work suggestion loop",[("O-04","Publish BEAS lowest cell"),("O-08","Flag stalled projects and the next gradient"),("SAL-12","Report lowest-conversion funnel step"),("CS-10","Report each client's next Fibonacci step"),("O-24","Turn signals into Work Proposal cards (outcome, falsifier, effort, owner)"),("O-T3","Check and rank by dependency order"),("O-23","Put proposals in the Owner Weekly Report"),("G-00","Approve / defer / decline"),("O-25","Approved → board, dispatched to team lead"),("O-16","After delivery: compare actual vs expected outcome"),("O-10","Integrate the gap; feed the next scan")]),
 ("W19","Reporting ladder",[("Task agents","Daily status on the board"),("Team lead","Daily board update; weekly team report"),("Hub lead","Weekly hub report"),("O-25","Daily blocked-items list"),("O-23","Owner Daily Brief; Owner Weekly Report"),("G-00","Reads, decides gates, sends requests back")]),
 ("W15","Agent fleet change",[("HRC-06","Propose new/changed agent role"),("G-06","Decision rights + autonomy level"),("G-05","Alignment review; downgrade path; instrumentation"),("G-00","Approve deployment"),("TEC-05","Deploy with logging + fallback"),("G-04","Audit first runs")]),
]

# ── 5. Human gates ──
GATES = [
 ("HG-01","Any material decision (the choice itself)","O-16","Framework cannot supply the choice","v5.1 L172"),
 ("HG-02","Naming the outcome after an action (e.g., 'success', 'failure')","Any","Framework cannot supply the naming","v5.1 L174"),
 ("HG-03","Situation-specific measurements an agent cannot take","G-03","Framework cannot supply empirical measurements","v5.1 L176"),
 ("HG-04","Stance on a held-open variable / any stipulation","G-01","Stipulations are operator choices (v5.3)","v5.3 §2"),
 ("HG-05","Strategic Direct (intent) for hub-level builds","O-03","Direct begins from Third-Kingdom awareness","v5.1 L1488–1497"),
 ("HG-06","Sales conversation (Creation stage)","SAL-05","'The human magic happens here'","v5.1 L1603"),
 ("HG-07","Honest escalation / walk-away on an objection or negotiation","SAL-06; SAL-10","Eliminate or Escalate — the honest path","v5.1 L1637; L2589"),
 ("HG-08","Contract signature; final negotiated terms","SAL-08; SAL-10","Derived (commitment)","Derived"),
 ("HG-09","Price and offer changes","FIN-06; SAL-07","Derived (commercial authority)","Derived"),
 ("HG-10","Capital deployment; budget approval","FIN-01; FIN-02","Derived (capital authority)","Derived"),
 ("HG-11","Hire, compensation, employment decisions","HRC-01; HRC-02; HRC-04","Derived (employment law/ethics)","Derived"),
 ("HG-12","Crisis command locus","O-19","Crisis step 3 — centralize to a clear Third-Kingdom locus","v5.1 L3383–3402"),
 ("HG-13","Blueprint approval (client + internal)","CS-03","'Written, documented, client-approved blueprint'","v5.1 L1339"),
 ("HG-14","Deploying or expanding any AI agent","G-05; TEC-05","AI Alignment — downgrade path before deployment","v5.1 L3933–3954"),
 ("HG-15","Legal positions and regulatory interpretation","HRC-10","Ambiguous rules held as bounded variables","v5.1 L3509–3526"),
 ("HG-16","Growth-stage classification for BEAS","O-04","Stage determines expected tilt","v5.1 L2486"),
 ("HG-17","Scale go/no-go","O-11","Patience — defer growth until the pillar is in place","v5.1 L2270–2287"),
 ("HG-18","Adopting a derived algorithm as a standing SOP","O-02","Derivations stay labelled; adoption is an operator decision","AGENTS.md rule 4"),
 ("HG-19","Publishing claims about results / the methodology","CON-*; MKT-*","Falsifiability + CC BY 4.0 attribution","v5.1 L1657; LICENSE"),
 ("HG-20","Refunds, credits, retention concessions","CS-05; CS-08","Derived (financial authority)","Derived"),
 ("HG-21","Psychological, health, family, grief, spiritual matters","G-00","Human-reserved algorithms — not delegated to agents","Derived (care boundary)"),
]

# ── 6. Held-open variables ──
HELD = [
 ("V-01","KPI thresholds (response times beyond 'instant', conversion targets, NPS target, DSO, churn %)","Canon names the metrics, not the numbers","Operator sets from own baseline measurements","Held open"),
 ("V-02","BEAS band calibration across industries / growth stages","v5.2 Run 4: 'exact band calibration across the 52 industries/stages (bounded)'","Track BEAS vs outcomes over cycles","Held open (bounded)"),
 ("V-03","Is the 9-department × 4-pillar decomposition substrate-dependent?","v5.2 Run 4 opened this variable","Test whether these nine hubs fit this specific business","Held open"),
 ("V-04","Does the co-creator proof hold for an AI or a collective?","v5.2 Run 5: held open — directly relevant to agents","Do not treat agents as co-creators; choice stays human (G-00)","Held open"),
 ("V-05","Autonomy level of each agent (A / H / L / R)","Derived in this map; canon does not assign autonomy","Operator adjusts per risk appetite; G-06 records","Stipulation"),
 ("V-06","Agent placement where canon names no hub (Risk, Innovation, Vendor, Compliance, Negotiation, Writing, Facilitation, Learning, Memory, Communication)","Derived placement","Re-home freely; responsibilities do not change","Stipulation"),
 ("V-07","Growth stage of the business (pre-revenue / scaling / mature)","Canon gives tilts per stage, not the classification","Founder classifies monthly (HG-16)","Held open"),
 ("V-08","Tool stack beyond those canon names (Lead Connector/GHL, Flows AI, n8n, Komodo, GFunnel University)","Not supplied","Operator selects; record in TEC-01 architecture doc","Held open"),
 ("V-09","Compensation numbers, prices other than the $297 downgrade path, budgets","Not supplied","Measured / decided by humans","Held open"),
 ("V-10","Legal jurisdiction and regulatory framework","Not supplied","HRC-10 maps once jurisdiction is known","Held open"),
 ("V-11","Number of agent instances per role (e.g., several SAL-01 instances)","Not supplied","Size from volume measurements","Held open"),
 ("V-12","Phase routing (Acquire → Scientific Inquiry, Create → Master Meta-Algorithm)","Proposed in substrate/2026-09-28-ace-phasing.md — substrate, NOT canon","O-02 may apply it as a labelled proposal","Live proposal (not canon)"),
 ("V-13","Whether 'Gather before Create' slows time-critical work (e.g., crisis)","Canon resolves by Crisis algorithm: triage first","O-19 overrides sequencing in confirmed crises only","Structurally derived"),
 ("V-14","v5.4 constructs (Line-Field, Knot-Line, Lamina) in business operations","v5.4 states zero novel predictions; not physics; not business ops","Excluded from this map","Excluded"),
]

# ═══════════════════════ WORKBOOK ═══════════════════════
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864"); H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=9); BOLD = Font(name=F, size=9, bold=True)
TITLE = Font(name=F, size=14, bold=True, color="1F3864")
INPUT = PatternFill("solid", fgColor="FFFF00")
TIER_FILL = {"0":"E2EFDA","1":"DDEBF7","2":"FFF2CC","3":"FCE4D6","4":"FFFFFF"}
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, row, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=3)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"

def body(ws, r, vals, fill=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
        if fill: cell.fill = PatternFill("solid", fgColor=fill)

wb = Workbook()
readme = wb.active; readme.title = "README"
org = wb.create_sheet("Org Chart")
deliv = wb.create_sheet("Deliverables")
rhy = wb.create_sheet("Operating Rhythm")
roster = wb.create_sheet("Agent Roster")
resp = wb.create_sheet("Responsibilities")
algc = wb.create_sheet("Algorithm Coverage")
elc = wb.create_sheet("Canon Element Coverage")
beas = wb.create_sheet("BEAS Accountability")
wfs = wb.create_sheet("Workflows")
hg = wb.create_sheet("Human Gates")
ho = wb.create_sheet("Held-Open Variables")
gap = wb.create_sheet("Gap Check")

# ── Responsibilities rows (built first so counts are known) ──
R = []  # (rid, agent, element_key, responsibility, source_type, source, lines, cadence, trigger, output, handoff, gate, status)
n = 0
alg_rows = []
for alg in ALGS:
    owner, cat, status, cadence = lookup(alg["name"])
    steps = alg["steps"]
    for (k, txt, ln) in steps:
        if isinstance(owner, dict):
            ag_ = owner.get(str(k), owner.get("*"))
        else:
            ag_ = owner
        n += 1
        nxt = f"Step {k+1}" if k < len(steps) else "Document & restart at higher baseline (Repetition)"
        gate = "Human-only" if ag_ == "G-00" else ""
        st = "Human-reserved" if ag_ == "G-00" and "Human-reserved" in status else ("Canon algorithm step" if "Derived" not in status else "Derived")
        R.append((f"R-{n:04d}", ag_, f"ALG:{alg['name']}", f"Step {k}: {txt}", "Algorithm step", alg["name"], f"v5.1 L{ln}", cadence, f"Step {k-1} complete" if k > 1 else "Algorithm triggered", nxt, ag_ if k < len(steps) else "O-15 (Database)", gate, st))
    alg_rows.append((alg["name"], cat, len(steps), f"v5.1 L{alg['start']}–{steps[-1][2]}", owner if isinstance(owner, str) else owner.get("*", ", ".join(sorted(set(owner.values())))), ", ".join(sorted(set(owner.values()))) if isinstance(owner, dict) else owner, status, cadence))
for (key, ag_, txt, src, lines, cad, trig, out, hand, gate) in T:
    n += 1
    R.append((f"R-{n:04d}", ag_, key, txt, "Canon table / protocol", src, lines, cad, trig, out, hand, gate, "Canon table" if not src.startswith("v5.3") else "Canon (v5.3 meta-tier)"))

for (ag_, key, txt, src, lines_, cad, trig, out, hand, gate) in TEAM_R:
    n += 1
    R.append((f"R-{n:04d}", ag_, key, txt, "Team layer (derived)", src, lines_, cad, trig, out, hand, gate, "Derived (team layer)"))

# Every agent must own ≥1 responsibility: add a charter responsibility per agent (mission).
for a in AGENTS:
    n += 1
    R.append((f"R-{n:04d}", a[0], f"CHARTER:{a[0]}", "Charter — " + a[5], "Agent charter (derived)", a[6], a[16], "Continuous", a[7], a[8], a[9], a[12] if a[12] not in ("—", "") else "", "Derived role charter"))

# ── README ──
readme["A1"] = "GFunnel Agent Org Map — Task-Specific Agents for a Business Operation"; readme["A1"].font = TITLE
lines = [
 ("What this is", "Every role and every responsibility needed to run a business on the GFunnel Methodology (Omni Process), decomposed into task-specific AI agents (plus the human gates they must escalate to)."),
 ("Status", "DERIVATION — not canon. Built under AGENTS.md rules 4 & 7: the roles are derived from canon; nothing here is a primary algorithm of the methodology, and nothing here changes the base code DETECT → PROCESS → RESPOND."),
 ("Sources read", "v5.1 (full operational sections: Layers IV–VII, ◇, Master Meta-Algorithm, all per-layer, domain and extended algorithms, self-audit, cannot-supply), v5.2 (Deep-Lens protocol; Runs 4–5), v5.3 (Meta-Tier ⊙). v5.4 deliberately excluded: its constructs are not business operations and it states zero novel predictions."),
 ("How it was built", "build_agent_map.py parses all 77 algorithms (662 numbered steps) straight from versions/v5.1 so each step becomes an owned responsibility; operational tables (pipeline, script, I.A.C.E., Immersion, BEAS, Five Modes, pillars, self-audit, Forcing Test, Deep-Lens) were transcribed row by row. Re-run the script to regenerate."),
 ("Tabs", "Deliverables — 'I need X done': each deliverable's accountable manager agent and the team that does it.\nOrg Chart — the reporting hierarchy as an indented tree: level, solid reporting line, dotted line, direct reports, reporting path.\nAgent Roster — every agent with mission, DETECT/PROCESS/RESPOND, hand-offs, tools named in canon, KPIs, human gate, autonomy, source.\nResponsibilities — every atomic responsibility, one row each, with owner, source line, cadence, trigger, output, hand-off.\nAlgorithm Coverage — all 77 algorithms → owner, status, steps mapped vs steps in canon.\nCanon Element Coverage — every operational table row → mapped responsibility.\nBEAS Accountability — 45-cell scoring tool with the accountable agent per cell.\nWorkflows — 19 end-to-end hand-off chains (W16 newsletter issue; W17 Owner request → done; W18 work suggestion loop; W19 reporting ladder).\nOperating Rhythm — every report and channel to the Owner: cadence, from, to, contents.\nHuman Gates — what agents must never decide.\nHeld-Open Variables — what canon does not supply (never filled with assumption).\nGap Check — live formulas proving no responsibility is unowned and no agent is idle."),
 ("Tiers", "0 · Human authority / Governance (meta-tier)  ·  1 · Orchestration & diagnostics  ·  2 · Department leads (the Nine Hubs)  ·  3 · Team leads (accountable for deliverables, manage a team)  ·  4 · Task agents"),
 ("Autonomy codes (derived)", "\n".join(f"{k} = {v}" for k, v in AUTONOMY.items())),
 ("Status labels", "Canon algorithm step / Canon table = transcribed from canon · Derived = placement or charter chosen for this map · Human-reserved = no agent executes · Held open = canon does not supply the value."),
 ("Editable cells", "Yellow cells on 'BEAS Accountability' are inputs (0, 0.5 or 1; coordination 0–1). Everything else is reference; totals and checks are formulas."),
 ("Attribution (CC BY 4.0)", "Source: GFunnel Methodology (Omni Process) v5.1–v5.3, Cameron Garlick / GFunnel, https://github.com/GFunnel-Tech/methodology, CC BY 4.0. This workbook adapts (restructures) the material into an agent org map; the adaptation is not endorsed by the author."),
]
for i, (k, v) in enumerate(lines, 3):
    readme.cell(row=i, column=1, value=k).font = BOLD
    c = readme.cell(row=i, column=2, value=v); c.font = BODY; c.alignment = WRAP
readme.column_dimensions["A"].width = 24; readme.column_dimensions["B"].width = 120
r0 = len(lines) + 4
readme.cell(row=r0, column=1, value="Live totals").font = BOLD
tot = [("Agents", "=COUNTA('Agent Roster'!A:A)-1"), ("Responsibilities", "=COUNTA(Responsibilities!A:A)-1"),
       ("Algorithms covered", "=COUNTA('Algorithm Coverage'!A:A)-1"), ("Canon elements covered", "=COUNTA('Canon Element Coverage'!A:A)-1"),
       ("Gap Check status", "='Gap Check'!B14")]
for i, (k, f) in enumerate(tot, r0 + 1):
    readme.cell(row=i, column=1, value=k).font = BODY
    c = readme.cell(row=i, column=2, value=f); c.font = BOLD; c.alignment = Alignment(horizontal="left")

# ── Agent Roster ──
cols = ["Agent ID","Agent Name","Tier","Hub / Department","Reports To","Kingdom (3K)","Yin/Yang Phase","Mission","Primary Algorithm(s)","DETECT — Triggers","RESPOND — Outputs","Hands Off To","Tools Named in Canon","KPIs / Observation Signals","Human Gate","Failure Mode Guarded","Autonomy","Role Basis","Canon Source","# Responsibilities","Dotted Line To"]
header(roster, 1, cols, [9,34,16,22,10,9,16,60,36,34,34,20,20,32,26,30,9,28,24,10,30])
for i, a in enumerate(AGENTS, 2):
    d = DEPTS[a[3]]
    vals = [a[0], a[1], a[2], d[0], a[4], d[1], d[2], a[5], a[6], a[7], a[8], a[9], a[10], a[11], a[12], a[13], a[14], a[15], a[16], f"=COUNTIF(Responsibilities!$B:$B,A{i})", DOTTED.get(a[0], "—")]
    body(roster, i, vals, TIER_FILL[a[2][0]])
    roster.cell(row=i, column=1).font = BOLD
NROST = len(AGENTS) + 1
roster.conditional_formatting.add(f"T2:T{NROST}", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="FF9999")))

# ── Responsibilities ──
cols = ["Resp ID","Agent ID","Agent Name","Hub","Element Key","Responsibility","Source Type","Source (algorithm / table)","Canon Location","Cadence","Trigger (DETECT)","Output / Next (RESPOND)","Hand-off To","Human Gate","Status"]
header(resp, 1, cols, [9,9,28,18,26,80,18,30,14,18,22,24,16,16,18])
for i, r in enumerate(R, 2):
    rid, ag_, key, txt, stype, src, ln, cad, trig, out, hand, gate, st = r
    vals = [rid, ag_, f"=IFERROR(INDEX('Agent Roster'!$B:$B,MATCH(B{i},'Agent Roster'!$A:$A,0)),\"UNASSIGNED\")",
            f"=IFERROR(INDEX('Agent Roster'!$D:$D,MATCH(B{i},'Agent Roster'!$A:$A,0)),\"UNASSIGNED\")", key, txt, stype, src, ln, cad, trig, out, hand, gate, st]
    body(resp, i, vals)
NRESP = len(R) + 1

# ── Algorithm Coverage ──
cols = ["Algorithm","Category","Steps in Canon (parsed)","Steps Mapped","Unmapped","Canon Location","Primary Owner","All Owners","Coverage Status","Cadence"]
header(algc, 1, cols, [44,24,12,10,10,18,12,30,40,26])
for i, (nm, cat, ns, loc, prim, owners, status, cad) in enumerate(alg_rows, 2):
    vals = [nm, cat, ns, f"=COUNTIFS(Responsibilities!$H:$H,A{i},Responsibilities!$G:$G,\"Algorithm step\")", f"=C{i}-D{i}", loc, prim, owners, status, cad]
    body(algc, i, vals)
    algc.cell(row=i, column=3).comment = None
NALG = len(alg_rows) + 1
algc.conditional_formatting.add(f"E2:E{NALG}", CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor="FF9999")))
algc.cell(row=NALG + 2, column=1, value="'Steps in Canon' is parsed from versions/v5.1 by build_agent_map.py (numbered lines inside each algorithm section). Canon's Fast-Reference index lists 48 algorithm rows (its footer says 47); the Extended section adds 29 → 77, matching canon's '70+'.").font = Font(name=F, size=8, italic=True)

# ── Canon Element Coverage ──
cols = ["Element Key","Canon Element","Source","Canon Location","Owner Agent","Responsibilities Mapped","Covered?"]
header(elc, 1, cols, [26,70,30,16,10,12,10])
for i, (key, ag_, txt, src, lines_, *_r) in enumerate(T, 2):
    vals = [key, txt.split(" Guard against")[0][:200], src, lines_, ag_, f"=COUNTIF(Responsibilities!$E:$E,A{i})", f"=IF(F{i}>0,\"YES\",\"GAP\")"]
    body(elc, i, vals)
NEL = len(T) + 1
elc.conditional_formatting.add(f"G2:G{NEL}", CellIsRule(operator="equal", formula=['"GAP"'], fill=PatternFill("solid", fgColor="FF9999")))

# ── BEAS Accountability (scoring tool) ──
beas["A1"] = "BEAS — Business Equilibrium Adherence Score (45-point monthly assessment)"; beas["A1"].font = TITLE
beas["A2"] = "Yellow cells: enter 0 (absent), 0.5 (partial) or 1 (fully implemented); coordination 0–1. Row 16 is an illustrative EXAMPLE and is not counted. Source: v5.1 L1389–1446."; beas["A2"].font = Font(name=F, size=9, italic=True)
cols = ["Department","Accountable Lead","Correctness","Correctness Agent","Complexity","Complexity Agent","Patience","Patience Agent","Resilience","Resilience Agent","Dept Total /4","Coordination /1","Quadrant"]
for c, h in enumerate(cols, 1):
    cell = beas.cell(row=4, column=c, value=h); cell.font, cell.fill, cell.border, cell.alignment = H_FONT, H_FILL, BORDER, Alignment(wrap_text=True)
    beas.column_dimensions[get_column_letter(c)].width = [22,12,11,26,11,26,11,26,11,26,11,12,16][c-1]
PA = {  # per-hub pillar agents (who supplies the evidence for that cell)
 "STR":("STR-02","STR-03","STR-01","STR-05"),"MKT":("MKT-03","MKT-02","MKT-00","MKT-01"),"SAL":("SAL-09","SAL-07","SAL-12","SAL-03"),
 "OPS":("OPS-02","OPS-03","OPS-03","OPS-04"),"FIN":("FIN-04","FIN-03","FIN-02","FIN-01"),"HRC":("HRC-04","HRC-06","HRC-01","HRC-10"),
 "TEC":("OPS-02","TEC-04","TEC-01","TEC-08"),"CS":("CS-03","CS-04","CS-10","CS-08"),"CON":("CON-03","CON-01","CON-05","CON-08")}
QUAD = {"STR":"High Yin","MKT":"High Yang","SAL":"High Yang","OPS":"High Yang","FIN":"High Yin","HRC":"High Yin","TEC":"Bridging","CS":"Bridging","CON":"High Yin"}
dv = DataValidation(type="list", formula1='"0,0.5,1"', allow_blank=True); beas.add_data_validation(dv)
dvc = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True); beas.add_data_validation(dvc)
for j, code in enumerate(["STR","MKT","SAL","OPS","FIN","HRC","TEC","CS","CON"]):
    r = 5 + j
    vals = [DEPTS[code][0], code + "-00", None, PA[code][0], None, PA[code][1], None, PA[code][2], None, PA[code][3], f"=SUM(C{r},E{r},G{r},I{r})", None, QUAD[code]]
    for c, v in enumerate(vals, 1):
        cell = beas.cell(row=r, column=c, value=v); cell.font, cell.border = BODY, BORDER
    for c in (3, 5, 7, 9):
        beas.cell(row=r, column=c).fill = INPUT; dv.add(beas.cell(row=r, column=c))
    beas.cell(row=r, column=12).fill = INPUT; dvc.add(beas.cell(row=r, column=12))
beas["A14"] = "TOTALS"; beas["A14"].font = BOLD
for c, col in ((3,"C"),(5,"E"),(7,"G"),(9,"I"),(11,"K"),(12,"L")):
    beas.cell(row=14, column=c, value=f"=SUM({col}5:{col}13)").font = BOLD
ex = ["EXAMPLE (not counted)", "SAL-00", 1, "", 0.5, "", 0.5, "", 0, "", "=SUM(C16,E16,G16,I16)", 0.5, "High Yang"]
for c, v in enumerate(ex, 1):
    cell = beas.cell(row=16, column=c, value=v); cell.font = Font(name=F, size=9, italic=True, color="808080")
summ = [("BEAS Total /45", "=K14+L14"),
        ("Band", '=IF(COUNT(C5:C13,E5:E13,G5:G13,I5:I13)=0,"Not yet scored",IF(B18>=40,"Dynamic Middle—Sustained",IF(B18>=32,"Equilibrium—Functional",IF(B18>=24,"Imbalance—Managed",IF(B18>=16,"Imbalance—Exposed","Collapse Risk")))))'),
        ("Canon action", '=IF(B19="Not yet scored","—",IF(B18>=40,"Scale intentionally. Begin next Fibonacci step. (O-04 → STR-03)",IF(B18>=32,"Identify lowest-scoring department. Apply Shepherd\'s Way. (O-04 → O-03)",IF(B18>=24,"Gap analysis. Prioritize by dependency order. (O-04 → O-03)",IF(B18>=16,"Return to Gather. Full Immersion Model recommended. (O-03)","Emergency blueprint session. Minimum viable docs in 9 depts within 30 days. (O-19 + G-00)")))))'),
        ("Lowest department total", "=IF(B19=\"Not yet scored\",\"—\",INDEX(A5:A13,MATCH(MIN(K5:K13),K5:K13,0)))"),
        ("High Yang quadrant avg /4", '=AVERAGE(K6:K8)'),
        ("High Yin quadrant avg /4", '=AVERAGE(K5,K9,K10,K13)'),
        ("Bridging quadrant avg /4", '=AVERAGE(K11:K12)')]
for i, (k, f) in enumerate(summ, 18):
    beas.cell(row=i, column=1, value=k).font = BOLD
    c = beas.cell(row=i, column=2, value=f); c.font = BOLD
    beas.merge_cells(start_row=i, start_column=2, end_row=i, end_column=10)
for r in (22, 23, 24):
    beas.cell(row=r, column=2).number_format = "0.00"
beas["A26"] = "Pillar evidence (what 'fully implemented' looks like) — v5.1 L1362–1367"; beas["A26"].font = BOLD
for i, (p, txt, own, l) in enumerate(PILL, 27):
    beas.cell(row=i, column=1, value=p).font = BOLD
    c = beas.cell(row=i, column=2, value=txt + f"  (Pillar steward: {own})"); c.font = BODY; c.alignment = WRAP
    beas.merge_cells(start_row=i, start_column=2, end_row=i, end_column=13); beas.row_dimensions[i].height = 30

# ── Workflows ──
cols = ["Workflow ID","Workflow","Seq","Agent ID","Agent Name","Action","Hands Off To"]
header(wfs, 1, cols, [10,36,6,10,32,70,12])
r = 2
for wid, name, steps in WF:
    for s, (aid, act) in enumerate(steps, 1):
        nxt = steps[s][0] if s < len(steps) else "— (cycle closes)"
        vals = [wid, name, s, aid, f"=IFERROR(INDEX('Agent Roster'!$B:$B,MATCH(D{r},'Agent Roster'!$A:$A,0)),\"(group / any)\")", act, nxt]
        body(wfs, r, vals, "F2F2F2" if int(wid[1:]) % 2 else None); r += 1

# ── Human Gates ──
cols = ["Gate ID","Decision Agents Must Escalate","Raised By","Basis","Canon Location","Decides"]
header(hg, 1, cols, [8,52,22,48,22,12])
for i, g in enumerate(GATES, 2):
    body(hg, i, list(g) + ["G-00"])

# ── Held-Open ──
cols = ["Var ID","Variable","Why Open (canon status)","How It Narrows / Who","Status"]
header(ho, 1, cols, [8,60,50,50,22])
for i, h in enumerate(HELD, 2):
    body(ho, i, list(h))

# ── Org Chart (hierarchy) ──
from collections import defaultdict
BY = {a[0]: a for a in AGENTS}
KIDS = defaultdict(list)
for a in AGENTS:
    KIDS[a[4]].append(a[0])
def unit(a):
    return UNITS.get(a[0]) or UNITS.get(a[3]) or DEPTS[a[3]][0]
org["A1"] = "Org Chart — reporting hierarchy (solid lines); dotted lines listed alongside"; org["A1"].font = TITLE
org["A2"] = "Derived structure: canon names the nine hubs and the layer algorithms, not who reports to whom. Redraw freely; responsibilities do not move with the boxes."; org["A2"].font = Font(name=F, size=9, italic=True)
cols = ["Level","Org Chart","Agent ID","Unit","Reports To","Dotted Line To","Autonomy","Direct Reports","Responsibilities","Reporting Path"]
header(org, 4, cols, [7,58,9,30,10,40,9,10,12,40])
org.freeze_panes = "C5"
ORDER = []
def walk(node, depth, path):
    a = BY[node]; ORDER.append((depth, node, path))
    kids = KIDS[node]
    # leaders before individual contributors; keep roster order otherwise
    kids = sorted(kids, key=lambda k: (0 if KIDS[k] else 1))
    for k in kids:
        walk(k, depth + 1, path + [node])
walk("G-00", 0, [])
assert len(ORDER) == len(AGENTS), "every agent must sit in the tree"
LEVEL_FILL = ["1F3864", "2F5597", "BDD7EE", "DDEBF7", "F2F2F2", "FFFFFF"]
for i, (depth, node, path) in enumerate(ORDER, 5):
    a = BY[node]
    label = ("    " * depth) + ("└─ " if depth else "") + a[1]
    vals = [depth, label, node, unit(a), a[4], DOTTED.get(node, "—"), a[14],
            f"=COUNTIF('Agent Roster'!$E:$E,C{i})", f"=COUNTIF(Responsibilities!$B:$B,C{i})",
            " → ".join(path + [node])]
    body(org, i, vals, LEVEL_FILL[min(depth, 5)] if depth > 1 else None)
    c = org.cell(row=i, column=2)
    if depth == 0:
        for col in range(1, len(cols) + 1):
            org.cell(row=i, column=col).fill = PatternFill("solid", fgColor=LEVEL_FILL[0]); org.cell(row=i, column=col).font = Font(name=F, size=10, bold=True, color="FFFFFF")
    elif depth == 1:
        for col in range(1, len(cols) + 1):
            org.cell(row=i, column=col).fill = PatternFill("solid", fgColor=LEVEL_FILL[1]); org.cell(row=i, column=col).font = Font(name=F, size=9, bold=True, color="FFFFFF")
    elif KIDS[node]:
        c.font = BOLD
NORG = len(ORDER) + 4
r2 = NORG + 2
org.cell(row=r2, column=2, value="Span of control (direct reports per manager)").font = BOLD
mgrs = [n for _, n, _ in ORDER if KIDS[n]]
for j, m in enumerate(mgrs, r2 + 1):
    org.cell(row=j, column=2, value=BY[m][1]).font = BODY
    org.cell(row=j, column=3, value=m).font = BODY
    org.cell(row=j, column=8, value=f"=COUNTIF('Agent Roster'!$E:$E,C{j})").font = BOLD

# ── Deliverables ──
deliv["A1"] = "Deliverables — 'I need X done: who owns it?'"; deliv["A1"].font = TITLE
deliv["A2"] = "Each deliverable has one accountable manager agent; the doers are the team that performs the steps. Ownership is derived (canon names functions, not deliverable owners)."; deliv["A2"].font = Font(name=F, size=9, italic=True)
cols = ["ID","Deliverable","Accountable Agent ID","Accountable Agent","Hub","Doers (team)","Algorithm / Process","Cadence","Human Gate"]
header(deliv, 4, cols, [8,44,11,36,22,52,34,22,24])
deliv.freeze_panes = "C5"
for i, (d_, acc, doers, alg, cad, gate) in enumerate(DELIV, 5):
    vals = [f"D-{i-4:03d}", d_, acc, f"=IFERROR(INDEX('Agent Roster'!$B:$B,MATCH(C{i},'Agent Roster'!$A:$A,0)),\"UNASSIGNED\")",
            f"=IFERROR(INDEX('Agent Roster'!$D:$D,MATCH(C{i},'Agent Roster'!$A:$A,0)),\"UNASSIGNED\")", doers, alg, cad, gate]
    body(deliv, i, vals)
NDEL = len(DELIV) + 4

# ── Operating Rhythm ──
rhy["A1"] = "Operating Rhythm — reporting and communication with the Owner"; rhy["A1"].font = TITLE
rhy["A2"] = "Everything reaches the Owner through one channel (O-23 Owner Liaison). Cadences are derived; exact send times and channels are held open for the Owner to set."; rhy["A2"].font = Font(name=F, size=9, italic=True)
cols = ["Cadence","Report / Channel","From","To","Contents","Canon anchor"]
header(rhy, 4, cols, [14,30,24,28,70,30])
rhy.freeze_panes = "C5"
for i, row in enumerate(RHYTHM, 5):
    body(rhy, i, list(row))

# ── Gap Check ──
gap["A1"] = "Gap Check — live proof of 'no gaps'"; gap["A1"].font = TITLE
checks = [
 ("Agents in roster", f"=COUNTA('Agent Roster'!A2:A{NROST})", "—"),
 ("Agents with zero responsibilities", f"=COUNTIF('Agent Roster'!T2:T{NROST},0)", "must be 0"),
 ("Responsibilities", f"=COUNTA(Responsibilities!A2:A{NRESP})", "—"),
 ("Responsibilities with no valid owner", f"=COUNTIF(Responsibilities!C2:C{NRESP},\"UNASSIGNED\")", "must be 0"),
 ("Algorithms in canon (parsed)", f"=COUNTA('Algorithm Coverage'!A2:A{NALG})", "77 expected"),
 ("Algorithm steps in canon", f"=SUM('Algorithm Coverage'!C2:C{NALG})", "—"),
 ("Algorithm steps unmapped", f"=SUM('Algorithm Coverage'!E2:E{NALG})", "must be 0"),
 ("Canon table elements uncovered", f"=COUNTIF('Canon Element Coverage'!G2:G{NEL},\"GAP\")", "must be 0"),
 ("Agents outside the org chart", f"=COUNTA('Agent Roster'!A2:A{NROST})-COUNTA('Org Chart'!C5:C{NORG})", "must be 0"),
 ("Deliverables with no valid accountable agent", f"=COUNTIF(Deliverables!D5:D{NDEL},\"UNASSIGNED\")", "must be 0"),
 ("Workflow steps with unknown agent", f"=COUNTIF(Workflows!E2:E{r-1},\"(group / any)\")", "group steps (Hub leads / Any agent) are intentional"),
]
for c, h in enumerate(["Check","Value","Rule"], 1):
    cell = gap.cell(row=2, column=c, value=h); cell.font, cell.fill = H_FONT, H_FILL
for i, (k, f, rule) in enumerate(checks, 3):
    gap.cell(row=i, column=1, value=k).font = BODY; gap.cell(row=i, column=2, value=f).font = BOLD; gap.cell(row=i, column=3, value=rule).font = BODY
gap["A14"] = "OVERALL"; gap["A12"].font = BOLD
gap["B14"] = '=IF(AND(B4=0,B6=0,B9=0,B10=0,B11=0,B12=0),"NO GAPS — every responsibility owned, every agent active, every canon step mapped","GAPS FOUND — see red cells")'
gap["B14"].font = Font(name=F, bold=True, size=11, color="006100")
gap.column_dimensions["A"].width = 40; gap.column_dimensions["B"].width = 70; gap.column_dimensions["C"].width = 50
gap.cell(row=16, column=1, value="Agents by tier").font = BOLD
for i, tier in enumerate(["0","1","2","3","4"], 17):
    gap.cell(row=i, column=1, value=f"Tier {tier}").font = BODY
    gap.cell(row=i, column=2, value=f"=COUNTIF('Agent Roster'!C2:C{NROST},\"{tier}*\")").font = BODY
gap.cell(row=23, column=1, value="Agents by autonomy").font = BOLD
for i, (k, v) in enumerate(AUTONOMY.items(), 24):
    gap.cell(row=i, column=1, value=v).font = BODY
    gap.cell(row=i, column=2, value=f"=COUNTIF('Agent Roster'!Q2:Q{NROST},\"{k}\")").font = BODY
gap.cell(row=29, column=1, value="Responsibilities per hub").font = BOLD
for i, code in enumerate(DEPTS, 30):
    gap.cell(row=i, column=1, value=DEPTS[code][0]).font = BODY
    gap.cell(row=i, column=2, value=f"=COUNTIF(Responsibilities!D2:D{NRESP},A{i})").font = BODY

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT)
print(f"wrote {OUT}: {len(AGENTS)} agents, {len(R)} responsibilities, {len(ALGS)} algorithms, {sum(len(a['steps']) for a in ALGS)} steps, {len(T)} table elements")
